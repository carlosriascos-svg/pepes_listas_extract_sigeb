"""
Procesa carpetas de reportes tipo "REPORTE - <timestamp>" que contienen, cada una,
un PDF de "Declaración de Bienes y Rentas / Conflictos de Interés" y un Excel
FORMULARIO.xlsx con dos hojas (diccionario de datos + datos planos).

Por cada subcarpeta encontrada:
  1. Lee el PDF y extrae campos clave -> "<pdf>.extraido.json" y "<pdf>.extraido.csv"
     (una sola fila) dentro de la misma subcarpeta.
  2. Lee la segunda hoja del Excel (los datos planos del formulario) y la vuelca a
     "<excel>.hoja2.csv" dentro de la misma subcarpeta.

Al final:
  3. Consolida todos los "*.hoja2.csv" en "output/consolidado_excel.xlsx", agregando
     las columnas `archivo_excel` y `ruta_excel` para saber de dónde vino cada fila.
  4. Consolida todos los "*.extraido.csv" (provenientes de los PDF) en
     "output/consolidado_pdfs.xlsx", agregando `archivo_pdf` y `ruta_pdf`.

Uso:
    python procesar_reportes.py "REPORTE - 2026-09-09T184701.504"
    python procesar_reportes.py "."   # procesa todas las subcarpetas "REPORTE*" encontradas
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd
import pdfplumber

# ---------------------------------------------------------------------------
# Extracción de campos del PDF
# ---------------------------------------------------------------------------
# Los regex asumen el layout estándar del formato oficial (Ley 2013 de 2019).
# Si la entidad cambia el formato del PDF, estos patrones deben ajustarse.

CAMPO_PATTERNS = {
    "tipo_declaracion": r"Tipo de declaraci[oó]n\s+(.+?)\s+Fecha de publicaci[oó]n",
    "fecha_publicacion": r"Fecha de publicaci[oó]n\s+([\d\-: ]+)",
    "tipo_documento": r"Tipo\s+(.+?)\s+N[uú]mero\s+(\d+)",
    "entidad": r"contratos y\s+(.+?)\s+ejecuten bienes",
    "direccion": r"Direcci[oó]n\s+(\[.*?\]|\S.*)",
    "cargo": r"Cargo o funci[oó]n que cumple\s+(.+)",
    "total_ingresos": r"TOTAL\s+\$?([\d.,]+)",
}


def _buscar(texto: str, patron: str, grupo: int = 1):
    m = re.search(patron, texto)
    if not m:
        return None
    return m.group(grupo).strip()


def _extraer_nombre_completo(texto: str) -> dict:
    m = re.search(
        r"Primer nombre Segundo nombre Primer apellido Segundo apellido\s*\n(.+)",
        texto,
    )
    if not m:
        return {
            "primer_nombre": None,
            "segundo_nombre": None,
            "primer_apellido": None,
            "segundo_apellido": None,
        }
    partes = m.group(1).split()
    # Formato típico: PRIMER_NOMBRE [SEGUNDO_NOMBRE] PRIMER_APELLIDO SEGUNDO_APELLIDO
    # Sin más contexto no es posible separar con certeza absoluta cuando hay nombres
    # compuestos; se asume 4 posiciones (nombre1, nombre2, apellido1, apellido2)
    # y se rellena con None si faltan.
    while len(partes) < 4:
        partes.insert(1, None)
    return {
        "primer_nombre": partes[0],
        "segundo_nombre": partes[1],
        "primer_apellido": partes[2],
        "segundo_apellido": partes[3],
    }


def _extraer_lugar(texto: str, encabezado: str) -> dict:
    patron = rf"{encabezado}\s*\nPa[ií]s\s+(\S+)\s+Departamento\s+(\S+)\s+Municipio\s+(\S+)"
    m = re.search(patron, texto)
    if not m:
        return {"pais": None, "departamento": None, "municipio": None}
    return {"pais": m.group(1), "departamento": m.group(2), "municipio": m.group(3)}


def extraer_campos_pdf(pdf_path: Path) -> dict:
    with pdfplumber.open(pdf_path) as pdf:
        texto = "\n".join(p.extract_text() or "" for p in pdf.pages)

    tipo_doc_match = re.search(r"Tipo\s+(.+?)\s+N[uú]mero\s+(\d+)", texto)
    nacimiento = _extraer_lugar(texto, "Lugar de nacimiento")
    domicilio = _extraer_lugar(texto, "Lugar de domicilio")
    nombres = _extraer_nombre_completo(texto)

    campos = {
        "archivo_pdf": pdf_path.name,
        "tipo_declaracion": _buscar(texto, CAMPO_PATTERNS["tipo_declaracion"]),
        "fecha_publicacion": _buscar(texto, CAMPO_PATTERNS["fecha_publicacion"]),
        **nombres,
        "tipo_documento": tipo_doc_match.group(1).strip() if tipo_doc_match else None,
        "numero_documento": tipo_doc_match.group(2).strip() if tipo_doc_match else None,
        "pais_nacimiento": nacimiento["pais"],
        "departamento_nacimiento": nacimiento["departamento"],
        "municipio_nacimiento": nacimiento["municipio"],
        "pais_domicilio": domicilio["pais"],
        "departamento_domicilio": domicilio["departamento"],
        "municipio_domicilio": domicilio["municipio"],
        "entidad": _buscar(texto, CAMPO_PATTERNS["entidad"]),
        "direccion": _buscar(texto, CAMPO_PATTERNS["direccion"]),
        "cargo": _buscar(texto, CAMPO_PATTERNS["cargo"]),
        "total_ingresos": _buscar(texto, CAMPO_PATTERNS["total_ingresos"]),
    }
    return campos


def guardar_json_csv(campos: dict, destino_base: Path) -> tuple[Path, Path]:
    json_path = destino_base.with_suffix(".extraido.json")
    csv_path = destino_base.with_suffix(".extraido.csv")

    json_path.write_text(
        json.dumps(campos, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(campos.keys()))
        writer.writeheader()
        writer.writerow(campos)

    return json_path, csv_path


# ---------------------------------------------------------------------------
# Extracción de la segunda hoja del Excel
# ---------------------------------------------------------------------------


def extraer_segunda_hoja_excel(xlsx_path: Path) -> Path:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if len(wb.sheetnames) < 2:
        raise ValueError(f"'{xlsx_path}' no tiene una segunda hoja.")

    hoja = wb[wb.sheetnames[1]]
    filas = list(hoja.iter_rows(values_only=True))

    csv_path = xlsx_path.with_suffix(".hoja2.csv")
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        for fila in filas:
            writer.writerow(fila)

    return csv_path


# ---------------------------------------------------------------------------
# Procesamiento de carpetas
# ---------------------------------------------------------------------------


def encontrar_carpetas_reporte(raiz: Path) -> list[Path]:
    if any(raiz.glob("*.pdf")) or any(raiz.glob("*.xlsx")):
        # La ruta dada ya es una carpeta de reporte individual
        return [raiz]
    return sorted(
        p
        for p in raiz.iterdir()
        if p.is_dir() and not p.name.startswith(".") and p.name.upper().startswith("REPORTE")
    )


def procesar_carpeta(carpeta: Path, errores: list[str]) -> dict:
    resultado = {"carpeta": str(carpeta), "pdf_csv": None, "excel_csv": None}

    pdfs = list(carpeta.glob("*.pdf"))
    xlsxs = list(carpeta.glob("*.xlsx"))

    if not pdfs:
        errores.append(f"[{carpeta}] No se encontró ningún .pdf")
    else:
        if len(pdfs) > 1:
            errores.append(
                f"[{carpeta}] Hay {len(pdfs)} PDFs, se usará solo '{pdfs[0].name}'"
            )
        try:
            campos = extraer_campos_pdf(pdfs[0])
            _, csv_path = guardar_json_csv(campos, pdfs[0])
            resultado["pdf_csv"] = csv_path
        except Exception as e:
            errores.append(f"[{carpeta}] Error leyendo PDF '{pdfs[0].name}': {e}")

    if not xlsxs:
        errores.append(f"[{carpeta}] No se encontró ningún .xlsx")
    else:
        if len(xlsxs) > 1:
            errores.append(
                f"[{carpeta}] Hay {len(xlsxs)} Excels, se usará solo '{xlsxs[0].name}'"
            )
        try:
            csv_path = extraer_segunda_hoja_excel(xlsxs[0])
            resultado["excel_csv"] = csv_path
        except Exception as e:
            errores.append(f"[{carpeta}] Error leyendo Excel '{xlsxs[0].name}': {e}")

    return resultado


# ---------------------------------------------------------------------------
# Consolidados
# ---------------------------------------------------------------------------


def consolidar_excel(resultados: list[dict], salida: Path) -> int:
    partes = []
    for r in resultados:
        csv_path = r["excel_csv"]
        if not csv_path:
            continue
        df = pd.read_csv(csv_path, header=0, dtype=str, encoding="utf-8-sig")
        df.insert(0, "ruta_excel", str(csv_path.parent))
        df.insert(0, "archivo_excel", csv_path.with_suffix("").with_suffix(".xlsx").name)
        partes.append(df)

    if not partes:
        return 0

    consolidado = pd.concat(partes, ignore_index=True)
    salida.parent.mkdir(parents=True, exist_ok=True)
    consolidado.to_excel(salida, index=False)
    return len(consolidado)


def consolidar_pdfs(resultados: list[dict], salida: Path) -> int:
    partes = []
    for r in resultados:
        csv_path = r["pdf_csv"]
        if not csv_path:
            continue
        df = pd.read_csv(csv_path, dtype=str, encoding="utf-8-sig")
        df.insert(0, "ruta_pdf", str(csv_path.parent))
        partes.append(df)

    if not partes:
        return 0

    consolidado = pd.concat(partes, ignore_index=True)
    salida.parent.mkdir(parents=True, exist_ok=True)
    consolidado.to_excel(salida, index=False)
    return len(consolidado)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Extrae campos de PDFs y la segunda hoja de Excels dentro de carpetas REPORTE."
    )
    parser.add_argument(
        "carpeta",
        type=str,
        help="Carpeta con las subcarpetas REPORTE (o una carpeta de reporte individual)",
    )
    parser.add_argument(
        "--output",
        type=str,
        default="output",
        help="Carpeta donde escribir los consolidados (por defecto: 'output')",
    )
    args = parser.parse_args()

    raiz = Path(args.carpeta)
    if not raiz.exists():
        raise NotADirectoryError(f"La ruta '{raiz}' no existe.")

    carpetas = encontrar_carpetas_reporte(raiz)
    print(f"Se encontraron {len(carpetas)} carpeta(s) de reporte.")

    errores: list[str] = []
    resultados = []
    for carpeta in carpetas:
        resultados.append(procesar_carpeta(carpeta, errores))

    output_dir = Path(args.output)
    n_excel = consolidar_excel(resultados, output_dir / "consolidado_excel.xlsx")
    n_pdf = consolidar_pdfs(resultados, output_dir / "consolidado_pdfs.xlsx")

    print(f"\nConsolidado Excel (hoja 2): {n_excel} fila(s) -> {output_dir / 'consolidado_excel.xlsx'}")
    print(f"Consolidado PDFs: {n_pdf} fila(s) -> {output_dir / 'consolidado_pdfs.xlsx'}")

    if errores:
        errores_path = output_dir / "errores.log"
        output_dir.mkdir(parents=True, exist_ok=True)
        errores_path.write_text("\n".join(errores), encoding="utf-8")
        print(f"\n{len(errores)} advertencia(s)/error(es). Ver detalle en {errores_path}")
    else:
        print("\nSin errores.")


if __name__ == "__main__":
    main()
